from docx import Document
from docx.shared import Pt

# Create a new Word document
doc = Document()

# Title
doc.add_heading('Resume', level=0)

# Contact Information
doc.add_heading('Contact Information', level=1)
doc.add_paragraph("Mobile: 6268442391, 7089704737")
doc.add_paragraph("Email: priteebaheshwar2023@gmail.com")
doc.add_paragraph("Address: AT+PO Bagholi, TH Lalburra, Dist. Balaghat, MP – 481441")

# Objective
doc.add_heading('Objective', level=1)
doc.add_paragraph(
    "TO OBTAIN A CHALLENGING CAREER IN THE GOVERNANCE AND PUT ALL MY EFFORTS "
    "INTO THE GROWTH OF THE ORGANISATION AND HAVE A GREAT WORKING ENVIRONMENT."
)

# Education Table
doc.add_heading('Education', level=1)
table = doc.add_table(rows=1, cols=4)
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Qualification'
hdr_cells[1].text = 'Board/University'
hdr_cells[2].text = 'Place'
hdr_cells[3].text = 'Year'

data = [
    ("10th", "MP Board", "Bhopal", "2011"),
    ("12th", "MP Board", "Bhopal", "2015"),
    ("GNM", "MPNRC", "Bhopal", "2017"),
    ("B.Sc", "MP Bhoj", "Bhopal", "2018"),
    ("PGDCA", "MCU", "Bhopal", "2024"),
]

for qual, board, place, year in data:
    row_cells = table.add_row().cells
    row_cells[0].text = qual
    row_cells[1].text = board
    row_cells[2].text = place
    row_cells[3].text = year

# Personal Details
doc.add_heading('Personal Details', level=1)
doc.add_paragraph("Name: Pritee Baheshwar")
doc.add_paragraph("Father’s Name: Ramdyal")
doc.add_paragraph("Mother’s Name: Nirmala")
doc.add_paragraph("Date of Birth: 13/04/1995")
doc.add_paragraph("Gender: Female")
doc.add_paragraph("Languages Known: Hindi, English")

# Experience
doc.add_heading('Experience', level=1)
doc.add_paragraph("Bal Gopal Children Hospital, Raipur")
doc.add_paragraph("Role: Nurse Trainee")
doc.add_paragraph("Duration: 05/10/2017 – 31/01/2022")

# Declaration
doc.add_heading('Declaration', level=1)
doc.add_paragraph("I hereby declare that the information provided above is true to the best of my knowledge.")
doc.add_paragraph("Date: __________")
doc.add_paragraph("Place: __________")

# Save the document
doc.save("Resume_Pritee_Baheshwar.docx")